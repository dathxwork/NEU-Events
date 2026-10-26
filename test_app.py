"""
Automated Test Suite for NEU Events Management System
Tests authentication, role permissions, discovery, registration, check-in, feedback, and moderation.
Compatible with the 5-user, 4-event compact dataset.
"""
import unittest
from datetime import date, timedelta
from app import app
from models import db, User, Event, EventCategory, Registration, AttendanceSession, AttendanceRecord, Feedback
from data_seeder import seed_database

class NeuEventsTestCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        app.config['TESTING'] = True
        app.config['WTF_CSRF_ENABLED'] = False
        seed_database(app)

    def setUp(self):
        self.client = app.test_client()

    def login(self, username, password):
        return self.client.post('/login', data={
            'username': username,
            'password': password
        }, follow_redirects=True)

    def logout(self):
        return self.client.get('/logout', follow_redirects=True)

    def test_01_login_success_student(self):
        """Test student login with valid credentials (11244035 / 123)"""
        response = self.login('11244035', '123')
        self.assertEqual(response.status_code, 200)
        self.assertIn('Nguyễn Thị Thu'.encode('utf-8'), response.data)
        self.logout()

    def test_02_login_invalid_password(self):
        """Test login fails with incorrect password"""
        response = self.login('11244035', 'wrongpass')
        self.assertEqual(response.status_code, 200)
        self.assertIn('không chính xác'.encode('utf-8'), response.data)

    def test_02b_login_invalid_mssv(self):
        """Test login fails when MSSV is not in NEU student registry"""
        response = self.login('99999999', 'anypassword')
        self.assertEqual(response.status_code, 200)
        self.assertIn('không tồn tại trong danh sách sinh viên Đại học Kinh tế Quốc dân'.encode('utf-8'), response.data)

    def test_02c_check_student_api(self):
        """Test public API for verifying MSSV in NEU database"""
        # Valid student
        res_valid = self.client.get('/api/check-student?student_id=11244035')
        self.assertEqual(res_valid.status_code, 200)
        json_valid = res_valid.get_json()
        self.assertTrue(json_valid['exists'])
        self.assertEqual(json_valid['student_id'], '11244035')
        self.assertEqual(json_valid['full_name'], 'Nguyễn Thị Thu')

        # Invalid student
        res_invalid = self.client.get('/api/check-student?student_id=99999999')
        self.assertEqual(res_invalid.status_code, 200)
        json_invalid = res_invalid.get_json()
        self.assertFalse(json_invalid['exists'])

    def test_03_event_discovery_page(self):
        """Test event listing and keyword search"""
        self.login('11244035', '123')
        response = self.client.get('/events')
        self.assertEqual(response.status_code, 200)
        self.assertIn('Khám phá sự kiện'.encode('utf-8'), response.data)
        self.assertIn(f'Ngày hội việc làm NEU {date.today().year}'.encode('utf-8'), response.data)
        self.logout()

    def test_04_event_detail_page(self):
        """Test viewing event detail for an existing event and verify dynamic status badge on banner"""
        self.login('11244035', '123')
        with app.app_context():
            job_event = Event.query.filter_by(code=f'NEU-JOB-{date.today().year}').first()
            job_id = job_event.id
            skill_event = Event.query.filter_by(code=f'NEU-SKILL-{date.today().year}').first()
            skill_id = skill_event.id
            hack_event = Event.query.filter_by(code=f'NEU-HACK-{date.today().year}').first()
            hack_id = hack_event.id

        # 1. Today's event shows "Đang diễn ra"
        response_job = self.client.get(f'/events/{job_id}')
        self.assertEqual(response_job.status_code, 200)
        self.assertIn(f'Ngày hội việc làm NEU {date.today().year}'.encode('utf-8'), response_job.data)
        self.assertIn('Đang diễn ra'.encode('utf-8'), response_job.data)

        # 2. Completed past event shows "Đã kết thúc"
        response_skill = self.client.get(f'/events/{skill_id}')
        self.assertEqual(response_skill.status_code, 200)
        self.assertIn('Đã kết thúc'.encode('utf-8'), response_skill.data)

        # 3. Future event shows "Sắp diễn ra"
        response_hack = self.client.get(f'/events/{hack_id}')
        self.assertEqual(response_hack.status_code, 200)
        self.assertIn('Sắp diễn ra'.encode('utf-8'), response_hack.data)

        self.logout()

    def test_05_my_events_page_student_only(self):
        """Test student's registered events schedule page is student-only"""
        # Student access -> OK
        self.login('11244035', '123')
        response = self.client.get('/my-events')
        self.assertEqual(response.status_code, 200)
        self.assertIn('Sự kiện đã đăng ký'.encode('utf-8'), response.data)
        self.logout()

        # Organizer access -> Blocked
        self.login('vieclam@neu.edu.vn', '123')
        org_resp = self.client.get('/my-events', follow_redirects=True)
        self.assertIn('chỉ dành riêng cho sinh viên'.encode('utf-8'), org_resp.data)
        self.logout()

    def test_06_checkin_page_student_view(self):
        """Test check-in page interface for student (QR-scan only and strictly registered events today)"""
        # Student 1 (Thu) has registered for JOB event happening today
        self.login('11244035', '123')
        response = self.client.get('/checkin')
        self.assertEqual(response.status_code, 200)
        self.assertIn('Điểm danh sự kiện'.encode('utf-8'), response.data)
        self.assertIn('Quét mã QR Code'.encode('utf-8'), response.data)
        self.assertNotIn('Quản lý phiên'.encode('utf-8'), response.data)
        # Dropdown only lists events student registered for today -> NEU-JOB is present, NEU-HACK is strictly absent
        self.assertIn(f'NEU-JOB-{date.today().year}'.encode('utf-8'), response.data)
        self.assertNotIn(f'NEU-HACK-{date.today().year}'.encode('utf-8'), response.data)
        self.logout()

        # Student 2 (Anh) has NO registered events today -> shows empty notice
        self.login('11240123', '123')
        res_empty = self.client.get('/checkin')
        self.assertEqual(res_empty.status_code, 200)
        self.assertIn('Không có sự kiện diễn ra hôm nay'.encode('utf-8'), res_empty.data)
        self.logout()

    def test_07_moderation_access_control(self):
        """Test RBAC: Student & Organizer cannot access moderation, but Manager can"""
        # Student access attempt -> Blocked
        self.login('11244035', '123')
        std_resp = self.client.get('/moderation', follow_redirects=True)
        self.assertEqual(std_resp.status_code, 200)
        self.assertIn('không có quyền'.encode('utf-8'), std_resp.data)
        self.logout()

        # Organizer access attempt -> Blocked
        self.login('vieclam@neu.edu.vn', '123')
        org_resp = self.client.get('/moderation', follow_redirects=True)
        self.assertEqual(org_resp.status_code, 200)
        self.assertIn('không có quyền'.encode('utf-8'), org_resp.data)
        self.logout()

        # Manager access attempt -> Allowed
        self.login('quanly@neu.edu.vn', '123')
        mgr_resp = self.client.get('/moderation')
        self.assertEqual(mgr_resp.status_code, 200)
        self.assertIn('Kiểm duyệt sự kiện'.encode('utf-8'), mgr_resp.data)
        self.assertIn('Hội thảo Ứng dụng AI'.encode('utf-8'), mgr_resp.data)
        self.logout()

    def test_08_organizer_events_page(self):
        """Test organizer events management page (strictly for ORGANIZER)"""
        # Organizer access -> Allowed
        self.login('vieclam@neu.edu.vn', '123')
        response = self.client.get('/organizer/events')
        self.assertEqual(response.status_code, 200)
        self.assertIn('Quản lý sự kiện'.encode('utf-8'), response.data)
        self.assertIn(f'Ngày hội việc làm NEU {date.today().year}'.encode('utf-8'), response.data)
        self.logout()

        # Manager access -> Blocked (Manager uses /moderation)
        self.login('quanly@neu.edu.vn', '123')
        mgr_resp = self.client.get('/organizer/events', follow_redirects=True)
        self.assertIn('chỉ dành riêng cho Ban tổ chức'.encode('utf-8'), mgr_resp.data)
        self.logout()

    def test_09_analytics_page(self):
        """Test reports and analytics page accessible by Manager & Organizer, blocked for Student"""
        # Manager -> OK
        self.login('quanly@neu.edu.vn', '123')
        response = self.client.get('/analytics')
        self.assertEqual(response.status_code, 200)
        self.assertIn('Báo cáo – thống kê'.encode('utf-8'), response.data)
        self.logout()

        # Student -> Blocked
        self.login('11244035', '123')
        std_resp = self.client.get('/analytics', follow_redirects=True)
        self.assertIn('không có quyền'.encode('utf-8'), std_resp.data)
        self.logout()

    def test_10_feedback_page_access(self):
        """Test feedback evaluation page accessible by Student and Organizer, blocked for Manager"""
        # Student -> OK
        self.login('11244035', '123')
        response = self.client.get('/feedback')
        self.assertEqual(response.status_code, 200)
        self.assertIn('Đánh giá sự kiện'.encode('utf-8'), response.data)
        self.logout()

        # Manager -> Blocked
        self.login('quanly@neu.edu.vn', '123')
        mgr_resp = self.client.get('/feedback', follow_redirects=True)
        self.assertIn('không có quyền truy cập chức năng đánh giá sự kiện'.encode('utf-8'), mgr_resp.data)
        self.logout()

    def test_11_manager_display_name(self):
        """Test manager login displays 'Đơn vị quản lý'"""
        login_res = self.login('quanly@neu.edu.vn', '123')
        self.assertEqual(login_res.status_code, 200)
        self.assertIn('Đơn vị quản lý'.encode('utf-8'), login_res.data)
        self.logout()

    def test_12_change_password_flow(self):
        """Test user changes password and can authenticate with new password"""
        self.login('vieclam@neu.edu.vn', '123')
        chg_res = self.client.post('/change-password', data={
            'old_password': '123',
            'new_password': 'organizerNewPass2026',
            'confirm_password': 'organizerNewPass2026'
        }, follow_redirects=True)
        self.assertEqual(chg_res.status_code, 200)
        self.assertIn('Đổi mật khẩu thành công'.encode('utf-8'), chg_res.data)
        self.logout()

        # Re-login with new password
        new_login = self.login('vieclam@neu.edu.vn', 'organizerNewPass2026')
        self.assertEqual(new_login.status_code, 200)
        self.logout()

        # Restore password for future tests
        with app.app_context():
            u = User.query.filter_by(username='vieclam@neu.edu.vn').first()
            if u:
                u.set_password('123')
                db.session.commit()

    def test_13_student_qr_checkin(self):
        """Test student self-check-in using QR Code token for active event"""
        self.login('11244035', '123')
        with app.app_context():
            event = Event.query.filter_by(code=f'NEU-JOB-{date.today().year}').first()
            session = AttendanceSession.query.filter_by(event_id=event.id, is_active=True).first()
            event_id = event.id
            pin = session.session_pin

        res = self.client.post('/checkin/submit', data={
            'event_id': str(event_id),
            'qr_code': pin
        }, follow_redirects=True)
        self.assertEqual(res.status_code, 200)
        self.assertIn('Điểm danh bằng mã QR thành công'.encode('utf-8'), res.data)

        # Verify record recorded with checkin_method = QR
        with app.app_context():
            u = User.query.filter_by(username='11244035').first()
            reg = Registration.query.filter_by(event_id=event_id, user_id=u.id).first()
            self.assertEqual(reg.attendance_status, 'ATTENDED')
            record = AttendanceRecord.query.filter_by(registration_id=reg.id).first()
            self.assertEqual(record.checkin_method, 'QR')
        self.logout()

    def test_13b_distinct_event_qr_codes_and_cross_rejection(self):
        """Test each event has distinct QR code and cross-event QR scanning is strictly rejected"""
        with app.app_context():
            e_job = Event.query.filter_by(code=f'NEU-JOB-{date.today().year}').first()
            e_hack = Event.query.filter_by(code=f'NEU-HACK-{date.today().year}').first()
            s_job = AttendanceSession.query.filter_by(event_id=e_job.id, is_active=True).first()
            s_hack = AttendanceSession.query.filter_by(event_id=e_hack.id, is_active=True).first()
            
            # Assert QR tokens and PINs are completely distinct across different events
            self.assertNotEqual(s_job.session_pin, s_hack.session_pin)
            self.assertNotEqual(s_job.qr_token, s_hack.qr_token)
            job_id = e_job.id
            hack_qr = s_hack.qr_token

        # Student logs in
        self.login('11240123', '123')

        # Try to check in to JOB event using HACK event's QR code -> Must be rejected
        cross_res = self.client.post('/checkin/submit', data={
            'event_id': str(job_id),
            'qr_code': hack_qr
        }, follow_redirects=True)
        self.assertEqual(cross_res.status_code, 200)
        self.assertIn('thuộc về sự kiện'.encode('utf-8'), cross_res.data)
        self.assertIn('Mỗi sự kiện có mã QR điểm danh riêng biệt'.encode('utf-8'), cross_res.data)
        self.logout()

    def test_14_student_cannot_manual_checkin_others(self):
        """Test student is forbidden from manually checking in other students"""
        self.login('11244035', '123')
        with app.app_context():
            reg = Registration.query.first()
            reg_id = reg.id

        res = self.client.post(f'/checkin/manual/{reg_id}', follow_redirects=True)
        self.assertEqual(res.status_code, 200)
        self.assertIn('Bạn không có quyền xác nhận điểm danh thủ công'.encode('utf-8'), res.data)
        self.logout()

    def test_15_role_based_navigation_isolation(self):
        """Test sidebar navigation strictly isolates Student, Organizer, and Manager"""
        # 1. Student navigation
        self.login('11244035', '123')
        std_events = self.client.get('/events')
        self.assertIn(b'href="/my-events"', std_events.data)
        self.assertIn(b'href="/checkin"', std_events.data)
        self.assertNotIn(b'href="/organizer/events"', std_events.data)
        self.assertNotIn(b'href="/moderation"', std_events.data)
        self.assertNotIn(b'href="/analytics"', std_events.data)
        self.logout()

        # 2. Organizer navigation
        self.login('vieclam@neu.edu.vn', '123')
        org_events = self.client.get('/events')
        self.assertIn(b'href="/organizer/events"', org_events.data)
        self.assertIn(b'href="/checkin"', org_events.data)
        self.assertIn(b'href="/feedback"', org_events.data)
        self.assertIn(b'href="/analytics"', org_events.data)
        self.assertNotIn(b'href="/my-events"', org_events.data)
        self.assertNotIn(b'href="/moderation"', org_events.data)
        self.logout()

        # 3. Manager navigation
        self.login('quanly@neu.edu.vn', '123')
        mgr_events = self.client.get('/events')
        self.assertIn(b'href="/moderation"', mgr_events.data)
        self.assertIn(b'href="/analytics"', mgr_events.data)
        self.assertNotIn(b'href="/my-events"', mgr_events.data)
        self.assertNotIn(b'href="/organizer/events"', mgr_events.data)
        self.assertNotIn(b'href="/checkin"', mgr_events.data)
        self.assertNotIn(b'href="/feedback"', mgr_events.data)
        self.logout()

    def test_16_organizer_cross_access_isolation(self):
        """Test Organizer 1 cannot edit or view feedback of events belonging to Organizer 2"""
        with app.app_context():
            it_event = Event.query.filter_by(code=f'NEU-HACK-{date.today().year}').first()
            it_event_id = it_event.id

        # Login as Organizer 1 (Career Center)
        self.login('vieclam@neu.edu.vn', '123')

        # Attempt to edit Organizer 2's event -> Blocked
        edit_res = self.client.get(f'/organizer/events?edit_id={it_event_id}', follow_redirects=True)
        self.assertIn('không có quyền truy cập hoặc chỉnh sửa sự kiện của Ban tổ chức khác'.encode('utf-8'), edit_res.data)

        # Attempt to view feedback of Organizer 2's event -> Blocked
        fb_res = self.client.get(f'/feedback?event_id={it_event_id}', follow_redirects=True)
        self.assertIn('không có quyền xem đánh giá của sự kiện do Ban tổ chức khác'.encode('utf-8'), fb_res.data)

        self.logout()

    def test_17_manager_moderation_workflow(self):
        """Test manager can approve a pending event"""
        self.login('quanly@neu.edu.vn', '123')
        with app.app_context():
            ai_event = Event.query.filter_by(code=f'NEU-AI-{date.today().year}').first()
            ai_event_id = ai_event.id
            self.assertEqual(ai_event.status, 'PENDING')

        # Manager approves event
        appr_res = self.client.post(f'/moderation/approve/{ai_event_id}', follow_redirects=True)
        self.assertEqual(appr_res.status_code, 200)
        self.assertIn('Đã phê duyệt sự kiện'.encode('utf-8'), appr_res.data)

        # Verify status is now APPROVED
        with app.app_context():
            ai_event = Event.query.get(ai_event_id)
            self.assertEqual(ai_event.status, 'APPROVED')

        self.logout()

    def test_18_cannot_cancel_attended_or_completed_event(self):
        """Test that attended or completed events cannot be cancelled by student or organizer"""
        # 1. Student attempts to cancel an event they already attended / completed
        self.login('11244035', '123')
        with app.app_context():
            skill_event = Event.query.filter_by(code=f'NEU-SKILL-{date.today().year}').first()
            user = User.query.filter_by(username='11244035').first()
            reg = Registration.query.filter_by(user_id=user.id, event_id=skill_event.id).first()
            reg_id = reg.id
            self.assertEqual(reg.attendance_status, 'ATTENDED')

        cancel_res = self.client.post(f'/registrations/{reg_id}/cancel', follow_redirects=True)
        self.assertEqual(cancel_res.status_code, 200)
        self.assertIn('Bạn đã điểm danh tham gia sự kiện này nên không thể hủy đăng ký!'.encode('utf-8'), cancel_res.data)

        # Verify status is still REGISTERED
        with app.app_context():
            reg_check = Registration.query.get(reg_id)
            self.assertEqual(reg_check.registration_status, 'REGISTERED')

        self.logout()

        # 2. Organizer attempts to cancel a completed event with attendees
        self.login('vieclam@neu.edu.vn', '123')
        with app.app_context():
            skill_event = Event.query.filter_by(code=f'NEU-SKILL-{date.today().year}').first()
            skill_id = skill_event.id
            self.assertEqual(skill_event.status, 'COMPLETED')

        org_cancel_res = self.client.post(f'/organizer/events/{skill_id}/cancel', follow_redirects=True)
        self.assertEqual(org_cancel_res.status_code, 200)
        self.assertIn('Sự kiện đã kết thúc nên không thể hủy!'.encode('utf-8'), org_cancel_res.data)

        # Verify event status is still COMPLETED
        with app.app_context():
            event_check = Event.query.get(skill_id)
            self.assertEqual(event_check.status, 'COMPLETED')

        self.logout()

    def test_22_ticket_api_and_qr_data(self):
        """Test API endpoint returns correct ticket and QR details"""
        with app.app_context():
            reg = Registration.query.filter_by(registration_status='REGISTERED').first()
            ticket_code = reg.ticket_code
            student_id = reg.user.student_id or reg.user.username

        res = self.client.get(f'/api/ticket/{ticket_code}')
        self.assertEqual(res.status_code, 200)
        data = res.get_json()
        self.assertTrue(data['exists'])
        self.assertEqual(data['ticket_code'], ticket_code)
        self.assertEqual(data['student_id'], student_id)

    def test_23_view_ticket_in_my_events(self):
        """Test student can load /my-events and view ticket triggers"""
        self.login('11244035', '123')
        res = self.client.get('/my-events')
        self.assertEqual(res.status_code, 200)
        html = res.get_data(as_text=True)
        self.assertIn('showTicketModal', html)
        self.assertIn('ticketModal', html)
        self.logout()

    def test_24_feedback_submission_and_system_time(self):
        """Test student submits and updates feedback with accurate system time"""
        from datetime import datetime, timedelta
        self.login('11240123', '123') # Student Minh Anh
        with app.app_context():
            skill_event = Event.query.filter_by(code=f'NEU-SKILL-{date.today().year}').first()
            skill_id = skill_event.id

        before_submit = datetime.now() - timedelta(seconds=2)
        res = self.client.post('/feedback/submit', data={
            'event_id': skill_id,
            'rating_stars': 5,
            'rating_content': 5,
            'rating_organization': 5,
            'rating_speaker': 5,
            'rating_venue': 5,
            'comment': 'Sự kiện rất tuyệt vời và bổ ích!',
            'is_anonymous': 'y'
        }, follow_redirects=True)
        after_submit = datetime.now() + timedelta(seconds=2)

        self.assertEqual(res.status_code, 200)
        with app.app_context():
            user_anh = User.query.filter_by(student_id='11240123').first()
            fb = Feedback.query.filter_by(event_id=skill_id, user_id=user_anh.id).first()
            self.assertIsNotNone(fb)
            self.assertTrue(before_submit <= fb.created_at <= after_submit)

        self.logout()

    def test_25_organizer_create_event_with_union_points_and_publication(self):
        """Test creating event with Youth Union points and publication, and verify points appear on banner"""
        from datetime import date, timedelta
        self.login('vieclam@neu.edu.vn', '123')
        
        # Test clicking create event opens new mode
        res_new = self.client.get('/organizer/events?new=1')
        self.assertEqual(res_new.status_code, 200)
        self.assertIn('Tạo mới sự kiện'.encode('utf-8'), res_new.data)

        # Submit new event with union points & publication
        today = date.today()
        res_save = self.client.post('/organizer/events/save', data={
            'title': 'Hội thảo Phát triển Kỹ năng Lãnh đạo 2026',
            'category_id': 1,
            'organizer_name': 'TT Tư vấn hướng nghiệp & Việc làm',
            'event_date': (today + timedelta(days=10)).strftime('%Y-%m-%d'),
            'start_time': '09:00',
            'end_time': '11:30',
            'location': 'Hội trường A2',
            'max_seats': 250,
            'reg_deadline': (today + timedelta(days=8)).strftime('%Y-%m-%d'),
            'fee': 'Miễn phí',
            'description': 'Chương trình rèn luyện kỹ năng lãnh đạo dành cho sinh viên NEU.',
            'union_points_plus': 12,
            'union_points_minus': 3,
            'handbook_name': 'An_pham_Lanh_dao_2026.pdf'
        }, follow_redirects=True)
        self.assertEqual(res_save.status_code, 200)

        with app.app_context():
            created_event = Event.query.filter_by(title='Hội thảo Phát triển Kỹ năng Lãnh đạo 2026').first()
            self.assertIsNotNone(created_event)
            self.assertEqual(created_event.union_points_plus, 12)
            self.assertEqual(created_event.union_points_minus, 3)
            self.assertEqual(created_event.handbook_name, 'An_pham_Lanh_dao_2026.pdf')
            evt_id = created_event.id

        # Verify points appear on event detail banner
        detail_res = self.client.get(f'/events/{evt_id}')
        self.assertEqual(detail_res.status_code, 200)
        self.assertIn('+12 điểm'.encode('utf-8'), detail_res.data)
        self.assertIn('Trừ 3 nếu vắng'.encode('utf-8'), detail_res.data)

        self.logout()

    def test_26_user_dropdown_and_modals_on_all_pages(self):
        """Test top-right user dropdown menu, user info modal, and change password modal dismissibility"""
        # Login as student
        self.login('11244035', '123')
        res = self.client.get('/events')
        self.assertEqual(res.status_code, 200)
        html = res.data.decode('utf-8')

        # Check user dropdown elements
        self.assertIn('id="userMenuWrapper"', html)
        self.assertIn('id="userDropdownMenu"', html)
        self.assertIn('Thông tin cá nhân', html)
        self.assertIn('Đổi mật khẩu', html)
        self.assertIn('Đăng xuất', html)

        # Check User Info Modal
        self.assertIn('id="userInfoModal"', html)
        self.assertIn('Nguyễn Thị Thu', html)
        self.assertIn('11244035', html)
        self.assertIn('11244035@st.neu.edu.vn', html)
        self.assertIn('Khoa Công nghệ thông tin', html)

        # Check Change Password Modal, eye visibility toggle, and outside click dismiss
        self.assertIn('id="changePasswordModal"', html)
        self.assertIn('togglePasswordVisibility', html)
        self.assertIn('closeChangePasswordModal', html)
        self.assertIn('backdropModals', html)

        # Check Academic Period Dropdowns for Student (6 years starting 2024-2025, Thu/Xuan/He)
        self.assertIn('id="academicYearSelect"', html)
        self.assertIn('id="semesterSelect"', html)
        self.assertIn('2024-2025', html)
        self.assertIn('2029-2030', html)
        self.assertIn('Học kỳ Thu 2026', html)
        self.assertNotIn('Tất cả năm học', html)

        self.logout()

        # Login as organizer and check same base header with unrestricted period options
        self.login('vieclam@neu.edu.vn', '123')
        res_org = self.client.get('/organizer/events')
        self.assertEqual(res_org.status_code, 200)
        html_org = res_org.data.decode('utf-8')
        self.assertIn('id="userMenuWrapper"', html_org)
        self.assertIn('id="userInfoModal"', html_org)
        self.assertIn('id="changePasswordModal"', html_org)
        self.assertIn('id="academicYearSelect"', html_org)
        self.assertIn('id="semesterSelect"', html_org)
        self.assertIn('Tất cả năm học', html_org)
        self.assertIn('Tất cả học kỳ', html_org)
        self.logout()

    def test_27_event_detail_banner_time_location_union_points_and_sponsor_benefits(self):
        """Test event detail banner displays event date, time, location, Union points (no 'rèn luyện'), and sponsor benefits in detail view"""
        self.login('11244035', '123')
        with app.app_context():
            job_event = Event.query.filter_by(code=f'NEU-JOB-{date.today().year}').first()
            job_id = job_event.id

        res = self.client.get(f'/events/{job_id}')
        self.assertEqual(res.status_code, 200)
        html = res.data.decode('utf-8')

        # 1. Banner displays date, time range, and location
        self.assertIn(job_event.event_date.strftime('%d/%m/%Y'), html)
        self.assertIn(job_event.start_time.strftime('%H:%M'), html)
        self.assertIn(job_event.end_time.strftime('%H:%M'), html)
        self.assertIn('Hội trường A2', html)

        # 2. Points are strictly 'Điểm Đoàn', never 'Điểm rèn luyện'
        self.assertIn('Điểm Đoàn', html)
        self.assertNotIn('rèn luyện', html)

        # 3. Sponsor perks are shown in event detail body
        self.assertIn('Quyền lợi & Quà tặng từ Nhà tài trợ', html)
        self.assertIn('Bộ quà tặng Welcome Kit NEU', html)
        self.assertIn('Voucher giảm giá 50%', html)

        self.logout()

        # 4. Test organizer creating a new event with sponsor benefits
        self.login('vieclam@neu.edu.vn', '123')
        today = date.today()
        res_save = self.client.post('/organizer/events/save', data={
            'title': 'Ngày hội Khởi nghiệp & Đổi mới sáng tạo 2026',
            'category_id': 1,
            'organizer_name': 'TT Tư vấn hướng nghiệp & Việc làm',
            'event_date': (today + timedelta(days=20)).strftime('%Y-%m-%d'),
            'start_time': '08:30',
            'end_time': '16:30',
            'location': 'Sảnh nhà A1',
            'max_seats': 200,
            'reg_deadline': (today + timedelta(days=18)).strftime('%Y-%m-%d'),
            'fee': 'Miễn phí',
            'description': 'Ngày hội kết nối các dự án khởi nghiệp sinh viên cùng các quỹ đầu tư.',
            'union_points_plus': 8,
            'union_points_minus': 2,
            'sponsor_benefits': '- Áo thun sự kiện NEU Startup\n- Vé tham dự Networking Dinner độc quyền\n- Quà tặng bình giữ nhiệt cao cấp'
        }, follow_redirects=True)
        self.assertEqual(res_save.status_code, 200)

        with app.app_context():
            created_evt = Event.query.filter_by(title='Ngày hội Khởi nghiệp & Đổi mới sáng tạo 2026').first()
            self.assertIsNotNone(created_evt)
            self.assertIn('Áo thun sự kiện NEU Startup', created_evt.sponsor_benefits)
            self.assertEqual(len(created_evt.sponsor_benefits_list), 3)
            evt_id = created_evt.id

        # Verify new event detail displays sponsor perks and no "rèn luyện"
        detail_res = self.client.get(f'/events/{evt_id}')
        self.assertEqual(detail_res.status_code, 200)
        detail_html = detail_res.data.decode('utf-8')
        self.assertIn('Áo thun sự kiện NEU Startup', detail_html)
        self.assertIn('Vé tham dự Networking Dinner độc quyền', detail_html)
        self.assertNotIn('rèn luyện', detail_html)

        self.logout()

    def test_28_pagination_and_calendar_smooth_scroll(self):
        """Test accurate pagination summary, page 1 to page 2 navigation, and calendar scroll IntoView"""
        from app import get_period_date_range
        self.login('11244035', '123')

        with app.app_context():
            start_d, end_d = get_period_date_range('2026-2027', 'ALL')
            expected_total = Event.query.filter(
                Event.status.in_(['APPROVED', 'COMPLETED', 'CANCELLED']),
                Event.event_date >= start_d,
                Event.event_date <= end_d
            ).count()

        # 1. Page 1 for period 2026-2027
        res_p1 = self.client.get('/events?academic_year=2026-2027&semester=ALL&page=1')
        self.assertEqual(res_p1.status_code, 200)
        html_p1 = res_p1.data.decode('utf-8')

        # Accurate summary: 1-6 of expected_total events
        self.assertIn(f'Hiển thị 1–6 trong tổng số {expected_total} sự kiện', html_p1)
        self.assertNotIn('trong tổng số 24 sự kiện', html_p1)

        # Pagination controls on Page 1: link to page 2 exists
        self.assertIn('page=2', html_p1)
        self.assertIn('Sau', html_p1)

        # Calendar toggle smooth scroll functions
        self.assertIn('scrollIntoView', html_p1)
        self.assertIn('scroll-mt-24', html_p1)

        # 2. Page 2 for 2026-2027
        res_p2 = self.client.get('/events?academic_year=2026-2027&semester=ALL&page=2')
        self.assertEqual(res_p2.status_code, 200)
        html_p2 = res_p2.data.decode('utf-8')

        # Accurate summary: 7-expected_total of expected_total events
        self.assertIn(f'Hiển thị 7–{expected_total} trong tổng số {expected_total} sự kiện', html_p2)
        # Link to page 1 exists
        self.assertIn('page=1', html_p2)
        self.assertIn('Trước', html_p2)

        self.logout()

    def test_29_union_points_badge_on_banner_and_footer_modals(self):
        """Test Union points badge displayed on bottom-right of event banner cards and footer modals for Terms & Privacy"""
        self.login('11244035', '123')
        res = self.client.get('/events')
        self.assertEqual(res.status_code, 200)
        html = res.data.decode('utf-8')

        # 1. Union points badge on bottom-right of event banner
        self.assertIn('absolute bottom-3 right-3', html)
        self.assertIn('điểm Đoàn', html)
        self.assertIn('fa-award', html)

        # 2. Updated footer string
        self.assertIn('Đại học Kinh tế Quốc dân - Cổng thông tin sự kiện - NEU Events.', html)

        # 3. Terms of Service Modal and Trigger Button
        self.assertIn('openTermsModal()', html)
        self.assertIn('id="termsModal"', html)
        self.assertIn('Quy định sử dụng NEU Events', html)

        # 4. Privacy Policy Modal and Trigger Button
        self.assertIn('openPrivacyModal()', html)
        self.assertIn('id="privacyModal"', html)
        self.assertIn('Chính sách bảo mật thông tin', html)

        self.logout()

    def test_30_help_modal_role_based_segregation(self):
        """Test Help Modal content is strictly segregated by user role (Student vs Organizer vs Manager) with exclusive role content"""
        # 1. Student view
        self.login('11244035', '123')
        res_stu = self.client.get('/events')
        self.assertEqual(res_stu.status_code, 200)
        html_stu = res_stu.data.decode('utf-8')
        self.assertIn('helpTabContent-student', html_stu)
        self.assertIn('Quy trình tham gia sự kiện dành cho Sinh viên NEU', html_stu)
        self.assertIn('Tìm kiếm & Đăng ký sự kiện', html_stu)
        self.assertNotIn('helpTabContent-organizer', html_stu)
        self.assertNotIn('helpTabContent-manager', html_stu)
        self.assertNotIn('helpTabBtn-organizer', html_stu)
        self.assertNotIn('helpTabBtn-manager', html_stu)
        self.assertNotIn('Cẩm nang nghiệp vụ dành cho Ban Tổ Chức sự kiện', html_stu)
        self.assertNotIn('Nghiệp vụ Phê duyệt & Giám sát dành cho Đơn vị Quản lý', html_stu)
        self.logout()

        # 2. Organizer view
        self.login('vieclam@neu.edu.vn', '123')
        res_org = self.client.get('/organizer/events')
        self.assertEqual(res_org.status_code, 200)
        html_org = res_org.data.decode('utf-8')
        self.assertIn('helpTabContent-organizer', html_org)
        self.assertIn('Cẩm nang nghiệp vụ dành cho Ban Tổ Chức sự kiện', html_org)
        self.assertIn('Khởi tạo & Gửi duyệt sự kiện', html_org)
        self.assertNotIn('helpTabContent-student', html_org)
        self.assertNotIn('helpTabContent-manager', html_org)
        self.assertNotIn('helpTabBtn-student', html_org)
        self.assertNotIn('helpTabBtn-manager', html_org)
        self.assertNotIn('Quy trình tham gia sự kiện dành cho Sinh viên NEU', html_org)
        self.assertNotIn('Nghiệp vụ Phê duyệt & Giám sát dành cho Đơn vị Quản lý', html_org)
        self.logout()

        # 3. Manager view
        self.login('quanly@neu.edu.vn', '123')
        res_mgr = self.client.get('/moderation')
        self.assertEqual(res_mgr.status_code, 200)
        html_mgr = res_mgr.data.decode('utf-8')
        self.assertIn('helpTabContent-manager', html_mgr)
        self.assertIn('Nghiệp vụ Phê duyệt & Giám sát dành cho Đơn vị Quản lý', html_mgr)
        self.assertNotIn('helpTabContent-student', html_mgr)
        self.assertNotIn('helpTabContent-organizer', html_mgr)
        self.assertNotIn('Quy trình tham gia sự kiện dành cho Sinh viên NEU', html_mgr)
        self.assertNotIn('Cẩm nang nghiệp vụ dành cho Ban Tổ Chức sự kiện', html_mgr)
        self.logout()

    def test_31_sidebar_centered_role_badge_and_no_duplicate_password_change(self):
        """Test sidebar displays centered role badge instead of section title, and no duplicate password change button at bottom-left"""
        self.login('11244035', '123')
        res = self.client.get('/events')
        self.assertEqual(res.status_code, 200)
        html = res.data.decode('utf-8')

        # 1. Centered role badge in sidebar
        self.assertIn('flex justify-center text-center', html)
        self.assertIn('fa-graduation-cap', html)
        self.assertIn('Sinh viên', html)

        # 2. Section title "Không gian sinh viên" replaced by badge
        self.assertNotIn('Không gian sinh viên</div>', html)

        # 3. Bottom actions: No "Đổi mật khẩu" button at bottom-left
        self.assertNotIn('Cài đặt & Hỗ trợ', html)
        self.assertIn('Hỗ trợ & Trợ giúp', html)

        # 4. Change password exists in user dropdown menu
        self.assertIn('openChangePasswordModal()', html)

        self.logout()

    def test_32_theme_switcher_light_dark_nature(self):
        """Test Theme Switcher functionality supporting light, dark, and nature modes with icon badge at top-right of user icon"""
        self.login('11244035', '123')
        res = self.client.get('/events')
        self.assertEqual(res.status_code, 200)
        html = res.data.decode('utf-8')

        # 1. Theme head initialization and CSS styles
        self.assertIn('neu_theme', html)
        self.assertIn('data-theme', html)
        self.assertIn('[data-theme="dark"]', html)
        self.assertIn('[data-theme="nature"]', html)

        # 2. Mini badge at top-right of user avatar icon
        self.assertIn('id="themeAvatarBadge"', html)
        self.assertIn('id="themeAvatarIcon"', html)

        # 3. Dedicated Theme Switcher button at top-right
        self.assertIn('id="themeMenuBtn"', html)
        self.assertIn('toggleThemeDropdown', html)
        self.assertIn('id="themeDropdownMenu"', html)

        # 4. Three theme options: Light, Dark, Nature
        self.assertIn("setTheme('light')", html)
        self.assertIn("setTheme('dark')", html)
        self.assertIn("setTheme('nature')", html)
        self.assertIn('Giao diện Sáng (Light)', html)
        self.assertIn('Giao diện Tối (Dark)', html)
        self.assertIn('Tự nhiên (Nature)', html)

        # 5. JS Theme update functions
        self.assertIn('function setTheme(', html)
        self.assertIn('function updateThemeUI(', html)

        # 6. Verify response contains no-cache headers
        self.assertIn('no-cache', res.headers.get('Cache-Control', ''))

        self.logout()

    def test_33_ticket_modal_responsive_non_overflowing(self):
        """Test Ticket Modal has non-overflowing responsive layout with fixed header/footer and scrollable body within 92vh"""
        self.login('11244035', '123')
        res = self.client.get('/events')
        self.assertEqual(res.status_code, 200)
        html = res.data.decode('utf-8')

        # 1. Non-overflowing modal container with viewport constraints
        self.assertIn('id="ticketModal"', html)
        self.assertIn('max-h-[92vh]', html)
        self.assertIn('overflow-y-auto', html)
        self.assertIn('flex flex-col', html)

        # 2. Fixed top brand header (shrink-0)
        self.assertIn('VÉ THAM DỰ SỰ KIỆN ĐIỆN TỬ', html)
        self.assertIn('closeTicketModal()', html)

        # 3. Fixed bottom actions footer (shrink-0)
        self.assertIn('downloadTicketQr()', html)
        self.assertIn('Tải ảnh QR', html)
        self.assertIn('Đóng vé', html)

        # 4. QR Code container optimized
        self.assertIn('id="ticketQrContainer"', html)
        self.assertIn('id="modalTicketCode"', html)

        self.logout()

    def test_34_terms_modal_role_based_segregation(self):
        """Test Terms of Service Modal is strictly segregated by user role (Student vs Organizer vs Manager) with exclusive role content"""
        # 1. Student view
        self.login('11244035', '123')
        res_stu = self.client.get('/events')
        self.assertEqual(res_stu.status_code, 200)
        html_stu = res_stu.data.decode('utf-8')
        self.assertIn('termsTabContent-student', html_stu)
        self.assertIn('Quy định & Quyền lợi tham gia sự kiện dành cho Sinh viên NEU', html_stu)
        self.assertIn('Tài khoản chính chủ & Sử dụng vé sự kiện', html_stu)
        self.assertNotIn('termsTabContent-organizer', html_stu)
        self.assertNotIn('termsTabContent-manager', html_stu)
        self.assertNotIn('termsTabBtn-organizer', html_stu)
        self.assertNotIn('termsTabBtn-manager', html_stu)
        self.assertNotIn('Quy chế tổ chức sự kiện & Quản lý điểm danh dành cho Ban Tổ Chức & CLB', html_stu)
        self.assertNotIn('Quy định Thẩm định, Phê duyệt & Giám sát Hệ thống dành cho Đơn vị Quản lý', html_stu)
        self.logout()

        # 2. Organizer view
        self.login('vieclam@neu.edu.vn', '123')
        res_org = self.client.get('/organizer/events')
        self.assertEqual(res_org.status_code, 200)
        html_org = res_org.data.decode('utf-8')
        self.assertIn('termsTabContent-organizer', html_org)
        self.assertIn('Quy chế tổ chức sự kiện & Quản lý điểm danh dành cho Ban Tổ Chức & CLB', html_org)
        self.assertIn('Khởi tạo hồ sơ & Tiêu chuẩn thông tin sự kiện', html_org)
        self.assertNotIn('termsTabContent-student', html_org)
        self.assertNotIn('termsTabContent-manager', html_org)
        self.assertNotIn('termsTabBtn-student', html_org)
        self.assertNotIn('termsTabBtn-manager', html_org)
        self.assertNotIn('Quy định & Quyền lợi tham gia sự kiện dành cho Sinh viên NEU', html_org)
        self.assertNotIn('Quy định Thẩm định, Phê duyệt & Giám sát Hệ thống dành cho Đơn vị Quản lý', html_org)
        self.logout()

        # 3. Manager view
        self.login('quanly@neu.edu.vn', '123')
        res_mgr = self.client.get('/moderation')
        self.assertEqual(res_mgr.status_code, 200)
        html_mgr = res_mgr.data.decode('utf-8')
        self.assertIn('termsTabContent-manager', html_mgr)
        self.assertIn('Quy định Thẩm định, Phê duyệt & Giám sát Hệ thống dành cho Đơn vị Quản lý', html_mgr)
        self.assertNotIn('termsTabContent-student', html_mgr)
        self.assertNotIn('termsTabContent-organizer', html_mgr)
        self.assertNotIn('Quy định & Quyền lợi tham gia sự kiện dành cho Sinh viên NEU', html_mgr)
        self.assertNotIn('Quy chế tổ chức sự kiện & Quản lý điểm danh dành cho Ban Tổ Chức & CLB', html_mgr)
        self.logout()

    def test_35_notification_button_and_dropdown_all_pages(self):
        """Test Notification Bell Button and Dropdown component on all pages with role-specific items and APIs"""
        # 1. Student view
        self.login('11244035', '123')
        pages = ['/events', '/my-events', '/checkin', '/feedback']
        for p in pages:
            res = self.client.get(p)
            self.assertEqual(res.status_code, 200, f'Failed on {p}')
            html = res.data.decode('utf-8')
            self.assertIn('id="notifBtn"', html, f'Missing notifBtn on {p}')
            self.assertIn('id="notifDropdown"', html, f'Missing notifDropdown on {p}')
            self.assertIn('id="notifBadge"', html, f'Missing notifBadge on {p}')
            self.assertIn('toggleNotifDropdown', html, f'Missing toggleNotifDropdown on {p}')
            self.assertIn('markAllNotifsAsRead', html, f'Missing markAllNotifsAsRead on {p}')
            self.assertIn('filterNotifs', html, f'Missing filterNotifs on {p}')
            self.assertIn('clearAllNotifs', html, f'Missing clearAllNotifs on {p}')

        # Test Notifications API endpoint
        api_res = self.client.get('/api/notifications')
        self.assertEqual(api_res.status_code, 200)
        data = api_res.get_json()
        self.assertTrue(data['success'])
        self.assertIn('notifications', data)
        self.assertIn('unread_count', data)
        # Verify notification links point to specific details (e.g., /events/)
        event_links = [n['link'] for n in data['notifications'] if '/events/' in n.get('link', '')]
        self.assertTrue(len(event_links) > 0, "Expected notifications to link directly to event details")

        # Test Mark All Read API
        read_all_res = self.client.post('/api/notifications/read-all')
        self.assertEqual(read_all_res.status_code, 200)
        self.assertEqual(read_all_res.get_json()['unread_count'], 0)

        # Test Clear All API
        clear_res = self.client.post('/api/notifications/clear')
        self.assertEqual(clear_res.status_code, 200)
        self.logout()

        # 2. Organizer view
        self.login('vieclam@neu.edu.vn', '123')
        res_org = self.client.get('/organizer/events')
        self.assertEqual(res_org.status_code, 200)
        html_org = res_org.data.decode('utf-8')
        self.assertIn('id="notifBtn"', html_org)
        self.assertIn('id="notifDropdown"', html_org)
        self.assertIn('toggleNotifDropdown', html_org)
        self.logout()

        # 3. Manager view
        self.login('quanly@neu.edu.vn', '123')
        res_mgr = self.client.get('/moderation')
        self.assertEqual(res_mgr.status_code, 200)
        html_mgr = res_mgr.data.decode('utf-8')
        self.assertIn('id="notifBtn"', html_mgr)
        self.assertIn('id="notifDropdown"', html_mgr)
        self.assertIn('toggleNotifDropdown', html_mgr)
        self.logout()

    def test_36_organizer_cannot_edit_completed_event(self):
        """Test completed events cannot be edited by Organizer (UI locked, backend blocked)"""
        self.login('vieclam@neu.edu.vn', '123')
        # Find or create a completed event owned by this organizer
        with app.app_context():
            org_user = User.query.filter_by(email='vieclam@neu.edu.vn').first()
            comp_ev = Event.query.filter_by(organizer_id=org_user.id, status='COMPLETED').first()
            if not comp_ev:
                # Find past event or create one
                comp_ev = Event.query.filter(Event.organizer_id == org_user.id, Event.event_date < date.today()).first()
            if not comp_ev:
                cat = EventCategory.query.first()
                comp_ev = Event(
                    code='NEU-TEST-COMP',
                    title='Sự kiện Test Đã Kết Thúc',
                    category_id=cat.id,
                    organizer_id=org_user.id,
                    organizer_name=org_user.department,
                    location='Hội trường A2',
                    event_date=date(2025, 1, 1),
                    start_time=time(8, 0),
                    end_time=time(11, 0),
                    registration_deadline=datetime(2024, 12, 31),
                    max_seats=100,
                    status='COMPLETED',
                    description='Mô tả'
                )
                db.session.add(comp_ev)
                db.session.commit()
            comp_id = comp_ev.id

        # 1. Access organizer events with edit_id pointing to completed event
        res = self.client.get(f'/organizer/events?edit_id={comp_id}', follow_redirects=True)
        self.assertEqual(res.status_code, 200)
        html = res.data.decode('utf-8')
        self.assertIn('Sự kiện đã kết thúc'.encode('utf-8'), res.data)
        self.assertIn('Khóa sửa'.encode('utf-8'), res.data)
        self.assertIn('disabled', html)

        # 2. Attempt to POST changes to completed event -> Must be rejected by backend
        post_data = {
            'event_id': str(comp_id),
            'title': 'Cố tình sửa sự kiện đã kết thúc',
            'category_id': '1',
            'organizer_name': 'Trung tâm Việc làm',
            'event_date': '2025-01-01',
            'start_time': '08:00',
            'end_time': '11:00',
            'location': 'Hội trường A2',
            'max_seats': '200',
            'reg_deadline': '2024-12-31',
            'fee': 'Miễn phí',
            'description': 'Mô tả sửa đổi',
            'union_points_plus': '5',
            'union_points_minus': '0'
        }
        post_res = self.client.post('/organizer/events/save', data=post_data, follow_redirects=True)
        self.assertEqual(post_res.status_code, 200)
        self.assertIn('Sự kiện đã kết thúc, không thể chỉnh sửa thông tin.'.encode('utf-8'), post_res.data)

        # Verify event title remained unchanged in DB
        with app.app_context():
            reloaded = Event.query.get(comp_id)
            self.assertNotEqual(reloaded.title, 'Cố tình sửa sự kiện đã kết thúc')
        self.logout()

    def test_37_event_status_empty_state_cancelled(self):
        """Test Empty State display when filtering status (e.g. cancelled) with 0 events"""
        # 1. On main events page with tab=cancelled (logged in)
        self.login('11244035', '123')
        res = self.client.get('/events?tab=cancelled')
        self.assertEqual(res.status_code, 200)
        html = res.data.decode('utf-8')
        # Empty state element should be present
        self.assertIn('Không có sự kiện đã hủy', html)
        self.logout()

        # 2. On organizer events page: Empty State row & status filter buttons should be present
        self.login('vieclam@neu.edu.vn', '123')
        org_res = self.client.get('/organizer/events')
        self.assertEqual(org_res.status_code, 200)
        org_html = org_res.data.decode('utf-8')
        self.assertIn('id="emptyFilterRow"', org_html)
        self.assertIn('setUnitStatusFilter', org_html)
        self.assertIn('data-status-btn="Đã hủy"', org_html)
        self.logout()

    def test_38_edit_event_requires_admin_reapproval(self):
        """Test editing an approved event resets status to PENDING and requires Manager approval"""
        self.login('vieclam@neu.edu.vn', '123')
        # Find an upcoming approved event owned by this organizer
        with app.app_context():
            org_user = User.query.filter_by(email='vieclam@neu.edu.vn').first()
            target_ev = Event.query.filter(
                Event.organizer_id == org_user.id,
                Event.status == 'APPROVED',
                Event.event_date >= date.today()
            ).first()
            if not target_ev:
                cat = EventCategory.query.first()
                target_ev = Event(
                    code='NEU-TEST-APPR',
                    title='Sự kiện Chuẩn bị Sửa',
                    category_id=cat.id,
                    organizer_id=org_user.id,
                    organizer_name=org_user.department,
                    location='Hội trường A2',
                    event_date=date(2026, 12, 1),
                    start_time=time(9, 0),
                    end_time=time(11, 30),
                    registration_deadline=datetime(2026, 11, 30),
                    max_seats=150,
                    status='APPROVED',
                    description='Mô tả gốc'
                )
                db.session.add(target_ev)
                db.session.commit()
            t_id = target_ev.id

        # Organizer edits the event
        edit_data = {
            'event_id': str(t_id),
            'title': 'Sự kiện Đã Cập Nhật Thời Gian',
            'category_id': '1',
            'organizer_name': 'Trung tâm Tư vấn Việc làm',
            'event_date': '2026-12-05',
            'start_time': '08:30',
            'end_time': '11:30',
            'location': 'Phòng B101',
            'max_seats': '180',
            'reg_deadline': '2026-12-04',
            'fee': 'Miễn phí',
            'description': 'Mô tả mới đã sửa',
            'union_points_plus': '5',
            'union_points_minus': '0'
        }
        res = self.client.post('/organizer/events/save', data=edit_data, follow_redirects=True)
        self.assertEqual(res.status_code, 200)
        self.assertIn('Chờ duyệt'.encode('utf-8'), res.data)

        # Check DB: status must be PENDING now!
        with app.app_context():
            ev_check = Event.query.get(t_id)
            self.assertEqual(ev_check.status, 'PENDING')
            self.assertEqual(ev_check.title, 'Sự kiện Đã Cập Nhật Thời Gian')
        self.logout()

        # Manager logs in and approves it
        self.login('quanly@neu.edu.vn', '123')
        mod_res = self.client.get('/moderation')
        self.assertEqual(mod_res.status_code, 200)
        self.assertIn('Sự kiện Đã Cập Nhật Thời Gian'.encode('utf-8'), mod_res.data)

        appr_res = self.client.post(f'/moderation/approve/{t_id}', follow_redirects=True)
        self.assertEqual(appr_res.status_code, 200)

        # Verify status is now APPROVED
        with app.app_context():
            ev_after = Event.query.get(t_id)
            self.assertEqual(ev_after.status, 'APPROVED')
        self.logout()

    def test_39_export_attendance_excel_and_csv(self):
        """Test Export Attendance defaults to Excel (.xlsx) using openpyxl, with CSV fallback"""
        self.login('vieclam@neu.edu.vn', '123')
        with app.app_context():
            org_user = User.query.filter_by(email='vieclam@neu.edu.vn').first()
            ev = Event.query.filter_by(organizer_id=org_user.id).first()
            ev_id = ev.id

        # 1. Default Export -> Excel (.xlsx)
        res_xlsx = self.client.get(f'/checkin/export/{ev_id}')
        self.assertEqual(res_xlsx.status_code, 200)
        self.assertEqual(
            res_xlsx.content_type,
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        self.assertIn('filename=', res_xlsx.headers.get('Content-Disposition', ''))
        self.assertTrue(res_xlsx.headers.get('Content-Disposition', '').endswith('.xlsx"'))

        # Verify openpyxl can read the downloaded binary stream
        from io import BytesIO
        import openpyxl
        wb = openpyxl.load_workbook(BytesIO(res_xlsx.data))
        self.assertIn('Diem_danh', wb.sheetnames)
        ws = wb['Diem_danh']
        self.assertIn('DANH SÁCH SINH VIÊN', ws['A1'].value)

        # 2. Format=csv -> UTF-8 BOM CSV
        res_csv = self.client.get(f'/checkin/export/{ev_id}?format=csv')
        self.assertEqual(res_csv.status_code, 200)
        self.assertTrue(res_csv.content_type.startswith('text/csv'))
        # Must start with UTF-8 BOM
        self.assertTrue(res_csv.data.startswith(b'\xef\xbb\xbf'))
        self.logout()


    def test_40_student_feedback_history_and_realtime_status(self):
        """Test student feedback page displays evaluation history, pending vs evaluated, and real-time timestamps"""
        self.login('11244035', '123') # Student Thu
        # 1. Access feedback page
        res = self.client.get('/feedback')
        self.assertEqual(res.status_code, 200)
        html = res.data.decode('utf-8')
        
        # Verify tabs and history elements
        self.assertIn('Lịch sử đánh giá sự kiện', html)
        self.assertIn('Chưa đánh giá', html)
        self.assertIn('Đã đánh giá', html)
        self.assertIn('switchStudentFeedbackTab', html)
        self.assertIn('filterStudentFeedbackList', html)

        # Verify real-time timestamp display
        self.assertIn('Thời gian gửi:', html)

        # 2. Submit new or updated feedback and verify redirect and updated timestamp
        with app.app_context():
            user_thu = User.query.filter_by(student_id='11244035').first()
            # Find an attended registration for Thu
            reg = Registration.query.filter_by(user_id=user_thu.id, attendance_status='ATTENDED').first()
            self.assertIsNotNone(reg)
            test_ev_id = reg.event_id

        update_res = self.client.post('/feedback/submit', data={
            'event_id': str(test_ev_id),
            'rating_stars': 5,
            'rating_content': 5,
            'rating_organization': 4,
            'rating_speaker': 5,
            'rating_venue': 4,
            'comment': 'Đánh giá cập nhật kiểm thử thời gian thực!',
            'is_anonymous': 'y'
        }, follow_redirects=True)
        self.assertEqual(update_res.status_code, 200)
        up_html = update_res.data.decode('utf-8')
        self.assertTrue('thành công theo thời gian hệ thống' in up_html or 'Cảm ơn bạn đã gửi đánh giá' in up_html)
        self.assertIn('Đánh giá cập nhật kiểm thử thời gian thực!', up_html)

        self.logout()

    def test_41_analytics_export_options_excel_csv_pdf(self):
        """Test analytics export report multi-format options: Excel (.xlsx), CSV (.csv), and PDF print view"""
        # 1. Login as manager
        self.login('quanly@neu.edu.vn', '123')
        res_page = self.client.get('/analytics')
        self.assertEqual(res_page.status_code, 200)
        page_html = res_page.data.decode('utf-8')

        # Check dropdown UI elements and links for all formats
        self.assertIn('Xuất báo cáo', page_html)
        self.assertIn('Xuất file Excel (.xlsx)', page_html)
        self.assertIn('Xuất file CSV (.csv)', page_html)
        self.assertIn('Xuất bản in / PDF (.pdf)', page_html)
        self.assertIn("format=xlsx", page_html)
        self.assertIn("format=csv", page_html)
        self.assertIn("format=pdf", page_html)

        # 2. Test Excel (.xlsx) export
        res_xlsx = self.client.get('/analytics/export?format=xlsx')
        self.assertEqual(res_xlsx.status_code, 200)
        self.assertEqual(res_xlsx.mimetype, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('.xlsx', res_xlsx.headers.get('Content-Disposition', ''))
        self.assertTrue(len(res_xlsx.data) > 1000)

        # 3. Test CSV export
        res_csv = self.client.get('/analytics/export?format=csv')
        self.assertEqual(res_csv.status_code, 200)
        self.assertIn('text/csv', res_csv.mimetype)
        csv_text = res_csv.data.decode('utf-8-sig')
        self.assertIn('BÁO CÁO THỐNG KÊ', csv_text)
        self.assertIn('TỔNG QUAN CHỈ SỐ KPI', csv_text)
        self.assertIn('DANH SÁCH CHI TIẾT', csv_text)

        # 4. Test PDF print view
        res_pdf = self.client.get('/analytics/export?format=pdf')
        self.assertEqual(res_pdf.status_code, 200)
        pdf_html = res_pdf.data.decode('utf-8')
        self.assertIn('BÁO CÁO TỔNG HỢP HIỆU QUẢ HOẠT ĐỘNG SỰ KIỆN', pdf_html)
        self.assertIn('window.print()', pdf_html)
        self.assertIn('ĐẠI HỌC KINH TẾ QUỐC DÂN', pdf_html)

        self.logout()

        # 5. Student cannot export analytics report
        self.login('11244035', '123')
        res_stu = self.client.get('/analytics/export?format=xlsx', follow_redirects=True)
        self.assertEqual(res_stu.status_code, 200)
        self.assertIn('báo cáo thống kê', res_stu.data.decode('utf-8'))
        self.logout()

    def test_42_analytics_dynamic_charts_and_filtering(self):
        """Test analytics dynamic chart computation and multi-criteria filtering with real DB data"""
        self.login('quanly@neu.edu.vn', '123')

        # 1. Default analytics page load
        res = self.client.get('/analytics')
        self.assertEqual(res.status_code, 200)
        html = res.data.decode('utf-8')

        # Verify dynamic JS chart payloads are rendered
        self.assertIn('const barChartData =', html)
        self.assertIn('const categoryData =', html)
        self.assertIn('weeklyBarChart', html)
        self.assertIn('categoryDonutChart', html)
        self.assertIn('btnModeWeek', html)
        self.assertIn('btnModeMonth', html)
        self.assertIn('btnModeQuarter', html)
        self.assertIn('btnModeYear', html)
        self.assertIn('btnChartTypeBar', html)
        self.assertIn('btnChartTypeLine', html)
        self.assertIn('chartMonthSelect', html)
        self.assertIn('chartYearSelect', html)
        self.assertIn('chartDatePicker', html)
        self.assertIn('chartActivePeriodBadge', html)
        self.assertIn('id="barChartTitle">Đăng ký & tham dự</h3>', html)
        self.assertNotIn('id="barChartTitle">Đăng ký & tham dự theo tuần</h3>', html)

        # 2. Filter by Category
        with app.app_context():
            cat = EventCategory.query.first()
            cat_id = cat.id if cat else 1
            expected_cat_events = Event.query.filter_by(category_id=cat_id).count()

        res_cat = self.client.get(f'/analytics?category_id={cat_id}')
        self.assertEqual(res_cat.status_code, 200)
        html_cat = res_cat.data.decode('utf-8')
        # Check center count and legend
        self.assertIn(f'{expected_cat_events}</span>', html_cat)

        # 3. Filter by Periods: Week, Month, Quarter, Year
        res_week = self.client.get('/analytics?period=THIS_WEEK')
        self.assertEqual(res_week.status_code, 200)

        res_period = self.client.get('/analytics?period=THIS_MONTH')
        self.assertEqual(res_period.status_code, 200)

        res_quarter = self.client.get('/analytics?period=THIS_QUARTER')
        self.assertEqual(res_quarter.status_code, 200)

        res_year = self.client.get('/analytics?period=THIS_YEAR')
        self.assertEqual(res_year.status_code, 200)

        # 4. Filter with Export Excel and CSV
        res_exp = self.client.get(f'/analytics/export?format=xlsx&category_id={cat_id}')
        self.assertEqual(res_exp.status_code, 200)
        self.assertEqual(res_exp.mimetype, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        res_exp_csv = self.client.get('/analytics/export?format=csv&period=THIS_WEEK')
        self.assertEqual(res_exp_csv.status_code, 200)

        self.logout()


if __name__ == '__main__':
    unittest.main()
